"""Regenerate explicit, fictional public examples; no network/API or customer data.
Optional local tools: python-docx, openpyxl, reportlab. Artifacts are committed;
Netlify's normal Node-only build does not require these Python dependencies.
"""
from pathlib import Path
import csv
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'samples';OUT.mkdir(exist_ok=True)
NOTICE='PREPARED FICTIONAL EXAMPLE — NOT A CUSTOMER DOCUMENT OR LIVE AI RESULT.'
rfp=f'''{NOTICE}
Harbor Borough — illustrative IT support RFP SAMPLE-001
All organizations, requirements, and statements below are fictional.

Section 1 — Incident response
The supplier must acknowledge priority-one incidents within 30 minutes.

Section 2 — Reporting
The supplier must provide a monthly service report.

Section 3 — Certification
The supplier must hold a current ISO 27001 certificate.
'''
evidence=f'''{NOTICE}
Cedar Support — illustrative company operating note
All statements are fictional; this note is not a certification or delivery record.

Section A — Incident acknowledgement
The service desk acknowledges priority-one incidents within 30 minutes.

Section B — Service reporting
The service team prepares a monthly service report.

Section C — Certification evidence
No certification evidence is included in this sample.
'''
(OUT/'fictional-rfp.txt').write_text(rfp,encoding='utf-8');(OUT/'fictional-company-evidence.txt').write_text(evidence,encoding='utf-8')
headers=['Requirement ID','Buyer requirement','Mandatory','Requirement source','Candidate company evidence','Evidence reference','Review status','Reviewer action']
rows=[['R1','The supplier must acknowledge priority-one incidents within 30 minutes.','Yes','fictional-rfp.txt §1','The service desk acknowledges priority-one incidents within 30 minutes.','fictional-company-evidence.txt §A','Needs Review','Verify current operating scope and delivery capability.'],['R2','The supplier must provide a monthly service report.','Yes','fictional-rfp.txt §2','The service team prepares a monthly service report.','fictional-company-evidence.txt §B','Needs Review','Verify report scope and accountable owner.'],['R3','The supplier must hold a current ISO 27001 certificate.','Yes','fictional-rfp.txt §3','No certificate supplied.','fictional-company-evidence.txt §C','Needs Review','Evidence gap: obtain and verify a current certificate; do not assert certification.']]
with (OUT/'illustrative-compliance-matrix.csv').open('w',encoding='utf-8-sig',newline='') as f:
 csv.writer(f).writerows([['Sample notice',NOTICE],headers,*rows])
w=Workbook();sheet=w.active;sheet.title='Review matrix';sheet.append(headers)
for row in rows:sheet.append(row)
sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1D4ED8')
for row in sheet:
 for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
for column,width in zip('ABCDEFGH',[18,60,15,28,60,40,22,75]):sheet.column_dimensions[column].width=width
for n in range(2,5):sheet.row_dimensions[n].height=64
note=w.create_sheet('Read me');note.append(['Notice',NOTICE]);note.append(['Purpose','Prepared illustration; not an automated extraction benchmark or a customer case study.']);note.append(['Review','All three rows remain Needs Review. No certification evidence is supplied.']);note.append(['Current product','The compliance workspace exports CSV/XLSX. The DOCX/PDF downloads are prepared examples.']);note.column_dimensions['A'].width=22;note.column_dimensions['B'].width=110
for row in note:
 for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
w.save(OUT/'illustrative-compliance-matrix.xlsx')
sections=[('Executive summary','This illustrative draft addresses the fictional Harbor Borough IT support request using only the supplied Cedar Support operating note. Review all statements and resolve the certification gap before using a proposal.'),('Incident acknowledgement','The supplied operating note states: “The service desk acknowledges priority-one incidents within 30 minutes.” [Company operating note §A; buyer requirement: RFP §1.] A reviewer must verify the note is current and delivery capability matches the proposed scope.'),('Monthly reporting','The supplied operating note states: “The service team prepares a monthly service report.” [Company operating note §B; buyer requirement: RFP §2.] Confirm the report scope and accountable owner.'),('Open qualification — certification','EVIDENCE GAP: No current ISO 27001 certificate is supplied. Do not assert certification. Obtain and verify the certificate and its scope before addressing RFP §3. The supplied evidence note §C explicitly records the missing evidence.'),('Before submission','Confirm requirements, delivery commitments, evidence validity, pricing, attachments, and the buyer’s submission instructions. This draft has not been approved for submission. No price, award result, or customer success metric is provided.')]
d=Document();d.add_heading('Illustrative proposal draft',0);d.add_paragraph(NOTICE);d.add_paragraph('Fictional Harbor Borough / Cedar Support · SAMPLE-001 · DRAFT FOR HUMAN REVIEW')
for title,text in sections:d.add_heading(title,1);d.add_paragraph(text)
d.add_paragraph('Prepared sample download. General DOCX/PDF export and an in-app proposal editor are not current workspace capabilities.')
for section in d.sections:section.footer.paragraphs[0].text='FICTIONAL SAMPLE • DRAFT FOR REVIEW'
d.styles['Normal'].font.size=Pt(10);d.save(OUT/'illustrative-proposal-draft.docx')
styles=getSampleStyleSheet();story=[Paragraph('Illustrative proposal draft',styles['Title']),Paragraph(escape(NOTICE),styles['BodyText']),Spacer(1,12),Paragraph('Fictional Harbor Borough / Cedar Support · SAMPLE-001 · DRAFT FOR HUMAN REVIEW',styles['BodyText'])]
for title,text in sections:story.extend([Spacer(1,14),Paragraph(escape(title),styles['Heading2']),Paragraph(escape(text),styles['BodyText'])])
story.extend([Spacer(1,14),Paragraph('Prepared sample download. General DOCX/PDF export and an in-app proposal editor are not current workspace capabilities.',styles['BodyText'])])
def footer(canvas,doc):
 canvas.saveState();canvas.setFont('Helvetica',8);canvas.drawString(40,25,'FICTIONAL SAMPLE | DRAFT FOR REVIEW');canvas.drawRightString(555,25,str(doc.page));canvas.restoreState()
SimpleDocTemplate(str(OUT/'illustrative-proposal-draft.pdf'),title='TenderCrest AI fictional proposal sample',author='TenderCrest AI',leftMargin=40,rightMargin=40,topMargin=35,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
print('Generated six explicit fictional sample assets; no API calls made.')
